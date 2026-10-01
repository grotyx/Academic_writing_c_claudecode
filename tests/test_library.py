"""Personal library: saved Word styles and templates, team profile, writing style copied into new papers."""
import json
from pathlib import Path

import pytest

from manuwright import library, lifecycle

ENGINE = Path(__file__).resolve().parents[1]


@pytest.fixture(autouse=True)
def home(tmp_path, monkeypatch):
    monkeypatch.setenv('MANUWRIGHT_HOME', str(tmp_path / 'home'))
    return tmp_path / 'home'


def make_template(path):
    from docx import Document
    from docx.shared import Pt
    doc = Document()
    doc.styles['Heading 1'].font.size = Pt(17)
    doc.add_paragraph('template text that must not leak into manuscripts')
    doc.save(path)
    return path


def test_word_template_saved_once_used_in_any_paper(tmp_path, capsys):
    template = make_template(tmp_path / 'team.docx')
    assert library.main(ENGINE, ['docx', 'add', str(template), '--name', 'Team style']) == 0
    assert library.docx_styles()['Team-style']['template']
    paper = tmp_path / 'paper'
    lifecycle.init(ENGINE, [str(paper)])
    answers = iter(['15', '3'])  # journal: none; Word style 3 = my saved style (keep, default, saved, custom)
    assert lifecycle.target(ENGINE, ['--project', str(paper / 'project.json')], ask=lambda _: next(answers)) == 0
    block = json.loads((paper / 'project.json').read_text())['docx']
    assert block == {'reference': 'templates/Team-style.docx'} and (paper / 'templates' / 'Team-style.docx').is_file()
    from harness.build import docx_style, document, append_markdown
    style = docx_style({'docx': block}, paper)
    doc = document(style)
    append_markdown(doc, '# Introduction\n\nBody text.', style)
    texts = [p.text for p in doc.paragraphs]
    assert texts == ['Introduction', 'Body text.']  # template content dropped, its heading style used
    assert doc.paragraphs[0].style.name == 'Heading 1' and doc.paragraphs[0].style.font.size.pt == 17


def test_reference_counts_in_review_snapshot(tmp_path):
    from tests.test_harness import put
    from harness.project import snapshot
    root = tmp_path
    put(root / 'templates/t.docx', 'x')
    put(root / 'project.json', '{}')
    files = snapshot(root / 'project.json', {'docx': {'reference': 'templates/t.docx'}, 'artifacts': []})
    assert any(k.endswith('t.docx') for k in files)


def test_init_copies_profile_and_writing_style(tmp_path):
    source = tmp_path / 'Style'
    (source / 'own').mkdir(parents=True)
    (source / 'own' / 'park_2024.md').write_text('# anchor', encoding='utf-8')
    (source / 'PDF' / 'own').mkdir(parents=True)
    (source / 'PDF' / 'own' / 'park_2024.pdf').write_bytes(b'%PDF')
    (source / 'terminology.md').write_text('# Terms', encoding='utf-8')
    assert library.main(ENGINE, ['style', 'import', str(source)]) == 0
    assert library.main(ENGINE, ['profile']) == 0
    assert library.profile_path().read_text(encoding='utf-8').startswith('# Author')
    paper = tmp_path / 'paper'
    lifecycle.init(ENGINE, [str(paper)])
    assert (paper / 'profile' / 'authors.md').is_file() and (paper / 'Style' / 'own' / 'park_2024.md').is_file()
    assert not (paper / 'Style' / 'PDF').exists()  # sources stay in the library
    assert json.loads((paper / 'project.json').read_text())['terminology'] == 'Style/terminology.md'


def test_style_add_files_sources_by_kind(tmp_path, capsys):
    pdf = tmp_path / 'smith_2024.pdf'
    pdf.write_bytes(b'%PDF')
    assert library.main(ENGINE, ['style', 'add', str(pdf), '--kind', 'landmark']) == 0
    assert (library.writing() / 'PDF' / 'landmark' / 'smith_2024.pdf').is_file()
    assert 'register my writing style' in capsys.readouterr().out
    assert library.main(ENGINE, ['style', 'add', str(pdf), '--kind', 'nope']) == 2


def test_library_never_overwrites_without_replace(tmp_path, capsys):
    one, two = make_template(tmp_path / 'one.docx'), make_template(tmp_path / 'two.docx')
    assert library.main(ENGINE, ['docx', 'add', str(one), '--name', 'team']) == 0
    assert library.main(ENGINE, ['docx', 'add', str(two), '--name', 'team']) == 1
    assert 'already exists' in capsys.readouterr().err
    assert library.main(ENGINE, ['docx', 'add', str(two), '--name', 'team', '--replace']) == 0
    md = tmp_path / 'a.md'
    md.write_text('first', encoding='utf-8')
    library.main(ENGINE, ['style', 'add', str(md)])
    md.write_text('second', encoding='utf-8')
    library.main(ENGINE, ['style', 'add', str(md)])
    assert (library.writing() / 'own' / 'a.md').read_text() == 'first'
    library.main(ENGINE, ['style', 'add', str(md), '--replace'])
    assert (library.writing() / 'own' / 'a.md').read_text() == 'second'
    library.main(ENGINE, ['profile'])
    assert library.main(ENGINE, ['profile', '--import', str(md)]) == 1
