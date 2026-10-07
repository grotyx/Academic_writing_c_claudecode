"""Academic writing mode: section cards, learned style profile, prose checks and their hooks."""
from __future__ import annotations

import importlib.util
import json
import subprocess
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'scripts'))
import academic_style as acad  # noqa: E402


def load_hook(name):
    spec = importlib.util.spec_from_file_location(name, ROOT / 'scripts' / 'hooks' / f'{name}.py')
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


BAD = """Low back pain plays a crucial role in disability, highlighting the need for care. We don't know why.
Furthermore, it is worth noting that outcomes were good. Moreover, we believe this is **very** important!

10 patients were significantly better in group A."""


@pytest.mark.parametrize('section', acad.SECTIONS)
def test_every_card_exists_and_its_model_text_passes_its_own_checks(section):
    text = (acad.STYLE_DIR / f'{section}.md').read_text(encoding='utf-8')
    model = text.split('## Model paragraph', 1)[1].split('\n', 1)[1] if '## Model paragraph' in text else ''
    assert acad.prose_issues(model, section) == []
    assert section in acad.card(section)


def test_prose_checks_flag_ai_register_and_form():
    codes = {(sev, code) for sev, code, _, _ in acad.prose_issues(BAD, 'results')}
    assert {('high', 'AI_PHRASE'), ('high', 'ING_TAIL'), ('high', 'CONTRACTION'), ('high', 'BOLD'),
            ('medium', 'BELIEF'), ('medium', 'CONNECTIVE_RUN'), ('medium', 'EXCLAMATION'),
            ('medium', 'NUMERAL_START'), ('medium', 'SIGNIFICANT_NO_STATS')} <= codes
    lines = {code: line for _, code, line, _ in acad.prose_issues(BAD, 'results')}
    assert lines['NUMERAL_START'] == 4 and lines['CONTRACTION'] == 1
    clean = ('Leg pain decreased more after endoscopic decompression (mean difference 1.2; 95% CI 0.4 to 2.0; '
             '*p* = 0.004).\n\n**Statistical analysis.** Data were analysed with R.\n\n**Keywords:** a; b; c')
    assert acad.prose_issues(clean, 'results') == []
    assert acad.prose_issues('Age did not differ significantly between groups (Table 1).', 'results') == []


def test_strict_blocks_only_what_pre_ai_papers_never_wrote():
    def codes(text):
        return [(i[0], i[1]) for i in acad.prose_issues(text, 'discussion', 40)]
    assert codes('Collagen plays a key role in disc degeneration.') == [('medium', 'AI_PHRASE')]
    assert codes('This improves care, paving the way for trials.') == [('high', 'AI_PHRASE')]  # reported once
    assert codes('These findings offer valuable insights.') == [('medium', 'AI_PHRASE')]
    assert acad.prose_issues('1. Department of Orthopaedic Surgery, Seoul.', 'title', 40) == []
    from check_style import split_sentences
    assert split_sentences('Pain fell (Fig. 2) at 6 mo. after surgery. It was a demo. Next.') == [
        'Pain fell (Fig. 2) at 6 mo. after surgery.', 'It was a demo.', 'Next.']


def test_words_the_card_calls_never_are_must_fix_but_statistical_leverage_is_not():
    def codes(text):
        return [(i[0], i[1]) for i in acad.prose_issues(text, 'methods', 40)]
    assert codes('We leveraged a registry to study the realm of surgery.') == [('high', 'AI_WORD'), ('high', 'AI_WORD')]
    assert codes('Observations with high leverage and high-leverage points were examined.') == []
    assert codes('Leverage values were inspected.') == []
    assert codes('Early care is crucial.') == [('medium', 'AI_WORD')]
    # a must-fix word inside a 'consider' tail clause is still must-fix, and reported once
    assert codes('We used a registry, showcasing the realm of surgery.') == [('high', 'AI_WORD'), ('high', 'AI_WORD')]


def test_split_sections_keeps_structured_abstract_labels_and_stops_at_references():
    text = ('Title\nAbstract\nBackground\nStenosis is common.\nMethods\nWe did it.\nIntroduction\nIntro.\n'
            '2. Materials and Methods\nMethod text.\nResults\nResult text.\nDiscussion\nDiscussion text.\n'
            'Conclusions\nConclusion.\nReferences\n1. A ref.')
    assert acad.split_sections(text) == {
        'abstract': 'Stenosis is common.\nWe did it.', 'introduction': 'Intro.', 'methods': 'Method text.',
        'results': 'Result text.', 'discussion': 'Discussion text.', 'conclusion': 'Conclusion.'}


def write_corpus(folder):
    intro = ('Lumbar spinal stenosis is a common cause of disability in older adults. The optimal extent of '
             'decompression remains uncertain. Previous studies were retrospective and small. This study aimed to '
             'compare two techniques in a prospective cohort. ') * 3
    discussion = ('In this cohort, endoscopic decompression was associated with fewer reoperations. This finding is '
                  'consistent with previous reports. However, residual confounding cannot be excluded. Further trials '
                  'are needed to confirm the association. ') * 3
    for i in (1, 2):
        (folder / f'paper{i}.md').write_text(f'Abstract\n\nShort.\n\nIntroduction\n\n{intro}\n\nDiscussion\n\n'
                                             f'{discussion}\n\nReferences\n\n1. Ref.\n', encoding='utf-8')
    from docx import Document
    doc = Document()
    doc.add_heading('Introduction', 1)
    doc.add_paragraph(intro)
    doc.add_heading('Discussion', 1)
    doc.add_paragraph(discussion)
    doc.save(folder / 'paper3.docx')
    (folder / 'example_template.md').write_text('Introduction\n\nignored ' * 50, encoding='utf-8')


def test_learn_measures_a_corpus_and_the_card_shows_it(tmp_path, monkeypatch):
    corpus = tmp_path / 'corpus'
    corpus.mkdir()
    write_corpus(corpus)
    profile = acad.learn([corpus], acad.library_profile_dir())
    assert profile['documents'] == 3 and set(profile['sections']) == {'introduction', 'discussion'}
    assert 'example_template.md' not in profile['sources']
    phrases = [p for p, _ in profile['sections']['discussion']['phrases']]
    assert 'consistent with previous reports' in phrases
    assert not any(any(c.isdigit() for c in p) for p in phrases)
    card = acad.card('discussion')
    assert 'Your corpus (learned from 3 document(s)' in card and 'Model paragraphs from your corpus' in card
    assert 'Not learned yet' in acad.card('methods') or 'has no methods text' in acad.card('methods')
    # A paper's own Style/profile wins over the library's.
    paper = tmp_path / 'paper'
    acad.learn([corpus / 'paper1.md'], paper / 'Style' / 'profile')
    assert 'learned from 1 document(s)' in acad.card('discussion', paper)
    # The learned profile quotes the sources, so it is never copied into a paper.
    from manuwright import library
    (library.writing() / 'own').mkdir(parents=True, exist_ok=True)
    (library.writing() / 'own' / 'mine.md').write_text('anchor', encoding='utf-8')
    _, copied = library.copy_into_paper(tmp_path / 'new_paper')
    assert [c.replace('\\', '/') for c in copied] == ['Style/own/mine.md']


def test_learned_or_reference_percentile_sets_the_long_sentence_limit(tmp_path):
    folder = tmp_path / 'p'
    (folder / 'Style' / 'profile').mkdir(parents=True)
    big = {'documents': 3, 'sentences': 40}
    (folder / 'Style' / 'profile' / 'style_profile.json').write_text(json.dumps({'sections': {
        'methods': {**big, 'sentence_length': {'p95': 31}},
        'results': {**big, 'sentence_length': {'p95': 90}},
        'discussion': {'documents': 1, 'sentences': 5, 'passive_pct': 20, 'we_per_100w': 1.0, 'hedges_per_100w': 1.0,
                       'sentence_length': {'mean': 15.0, 'sd': 2.0, 'p90': 17, 'p95': 17}}}}))
    ref = acad.reference_profile()['overall']
    # A learned limit never drops below the journals' 90th percentile ...
    assert acad.long_limit_for('methods', folder) == ref['methods']['p90_len']
    assert acad.long_limit_for('results', folder) == 55
    # ... and a corpus of one paper is shown on the card but never enforced.
    assert acad.long_limit_for('discussion', folder) == ref['discussion']['p95_len']
    assert 'too small to enforce' in acad.card('discussion', folder)


def test_mode_comes_from_the_environment_then_config(monkeypatch):
    assert acad.mode() == 'academic'
    from manuwright import lifecycle
    assert lifecycle.mode(['strict']) == 0 and acad.mode() == 'strict'
    monkeypatch.setenv('MANUWRIGHT_WRITING_MODE', 'off')
    assert acad.mode() == 'off'
    assert lifecycle.mode(['loud']) == 2


def test_strict_mode_blocks_high_findings_in_new_manuscript_text(monkeypatch):
    gates = load_hook('enforce_gates')
    edit = {'tool_name': 'Edit', 'cwd': '/p', 'tool_input': {
        'file_path': '/p/drafts/revision/REV1/06_discussion_REV1.md', 'old_string': 'x',
        'new_string': 'This plays a pivotal role, highlighting the importance of early care.'}}
    assert gates.decide(edit) is None  # academic mode reports after the edit; it never blocks
    monkeypatch.setenv('MANUWRIGHT_WRITING_MODE', 'strict')
    reason = gates.decide(edit)
    assert 'BLOCKED by academic writing mode (strict)' in reason and 'ING_TAIL' in reason
    assert 'manuwright style card discussion' in reason
    edit['tool_input']['new_string'] = 'Endoscopic decompression was associated with fewer reoperations.'
    assert gates.decide(edit) is None
    patch = ('*** Begin Patch\n*** Update File: drafts/revision/notes.md\n+We don\'t know.\n'
             '*** Update File: drafts/revision/06_discussion.md\n@@\n-old\n+Fine text here.\n*** End Patch')
    assert gates.decide({'tool_name': 'apply_patch', 'cwd': '/p', 'tool_input': {'command': patch}}) is None  # notes: not prose
    patch = patch.replace("+We don\'t know.", '+Fine.').replace('+Fine text here.', "+We don\'t know.")
    assert gates.decide({'tool_name': 'apply_patch', 'cwd': '/p', 'tool_input': {'command': patch}}).count('CONTRACTION') == 1
    for name in ('08_references.md', 'revision/REV1/response_letter_REV1.md', 'notes.md'):
        ref = {'tool_name': 'Write', 'cwd': '/p', 'tool_input': {'file_path': f'/p/drafts/{name}', 'content': "It's key."}}
        assert gates.strict_style(ref, f'/p/drafts/{name}') is None  # strict covers sections only
    plan = {'tool_name': 'Write', 'cwd': '/p', 'tool_input': {'file_path': '/p/drafts/draft_plan.md',
                                                            'content': "We don't block plans."}}
    assert gates.decide(plan) is None


def test_lint_after_edit_reports_academic_findings(tmp_path, monkeypatch):
    lint = load_hook('lint_on_edit')
    target = tmp_path / 'drafts' / '06_discussion.md'
    target.parent.mkdir()
    target.write_text('This plays a pivotal role in care.\n', encoding='utf-8')
    event = {'tool_name': 'Write', 'cwd': str(tmp_path), 'tool_input': {'file_path': str(target)}}
    code, message = lint.evaluate(event)
    assert code == 2 and '[ACADEMIC/MUST FIX/AI_PHRASE]' in message and 'manuwright style card discussion' in message
    monkeypatch.setenv('MANUWRIGHT_WRITING_MODE', 'off')
    assert 'ACADEMIC' not in lint.evaluate(event)[1]


def test_prompt_hook_injects_the_card_for_a_drafting_request_and_ignores_talk_about_the_tool(monkeypatch):
    intent = load_hook('style_intent')
    question = ('이 프로그램도 학술적으로 글을 쓰는 모드 같은 거를 개선할 수 있는 방법이 있을까? '
                '글 자체를 학술적으로 쓸 수 있도록 바꿀 수 있는 방법')
    assert intent.evaluate({'prompt': question}) == ''
    assert intent.evaluate({'prompt': '서론 써줘'}) == ''  # not in a paper folder: no card
    out = intent.evaluate({'prompt': '서론 써줘', 'cwd': str(ROOT)})  # the template checkout is a paper folder
    assert 'ACADEMIC STYLE CARD: introduction' in out
    assert intent.requested_sections('결과 써줘') == ['results'] and intent.requested_sections('방법을 다시 써줘') == ['methods']
    assert intent.evaluate({'prompt': 'Write a unit test for the abstract class', 'cwd': '/tmp'}) == ''
    assert intent.requested_sections('drafts/06_discussion.md 다시 써줘, 결론도 작성') == ['discussion', 'conclusion']
    assert intent.requested_sections('결과 나오면 알려줘') == []
    assert intent.requested_sections('그림 제목 수정해줘') == [] and not intent.detect('학술 검색해서 정리해줘')
    for statement in ('I turned academic mode off', '현재 학술 모드는 off', '방금 학술 모드 꺼'):
        assert intent.mode_toggle(statement) is None
    assert intent.mode_toggle('please turn academic mode off') == 'off' and intent.mode_toggle('이제 학술 모드 꺼줘') == 'off'
    assert intent.requested_sections('Introduction を書いて') == ['introduction'] and intent.requested_sections('重写 Discussion') == ['discussion']
    assert intent.requested_sections('suggest titles') == ['title'] and intent.detect('make this paragraph academic')
    assert intent.requested_sections('제목 후보 몇 개 줘') == ['title'] and intent.requested_sections('결과 정리해줘') == ['results']
    monkeypatch.setenv('MANUWRIGHT_WRITING_MODE', 'off')
    assert intent.evaluate({'prompt': '서론 써줘', 'cwd': str(ROOT)}) == ''


def test_session_start_injects_the_core_card_unless_off(tmp_path, monkeypatch):
    script = ROOT / 'scripts' / 'hooks' / 'session_contract.py'
    run = lambda folder: subprocess.run([sys.executable, str(script), str(folder)], capture_output=True,  # noqa: E731
                                        text=True, encoding='utf-8').stdout
    paper = tmp_path / 'paper'
    paper.mkdir()
    (paper / 'project.json').write_text('{"artifacts": []}', encoding='utf-8')
    out = run(paper)
    assert 'ACADEMIC WRITING MODE' in out and 'manuwright style card' in out
    assert 'ACADEMIC WRITING MODE' not in run(tmp_path)  # another project: the plugin hook stays quiet
    monkeypatch.setenv('MANUWRIGHT_WRITING_MODE', 'off')
    out = run(paper)
    assert 'WORKFLOW CONTRACT' in out and 'ACADEMIC WRITING MODE' not in out


def test_cli_style_commands_and_check_row(tmp_path, capsys, monkeypatch):
    out = subprocess.run([sys.executable, '-m', 'manuwright.cli', 'style', 'card', 'results'], capture_output=True,
                         text=True, encoding='utf-8', cwd=str(ROOT)).stdout
    assert 'ACADEMIC STYLE CARD: results' in out
    bad = tmp_path / '05_results.md'
    bad.write_text("We don't know.\n", encoding='utf-8')
    assert acad.main(['check', str(bad)]) == 1
    from manuwright import lifecycle
    monkeypatch.setattr(lifecycle, 'latest_release', lambda: None)
    monkeypatch.setattr(lifecycle, 'adapter_checks', lambda engine, current: [])
    lifecycle.check(lifecycle.Path(ROOT), [])
    report = capsys.readouterr().out
    assert '✓ Writing mode' in report and '· Learned style' in report


def test_a_pdf_without_a_reader_is_skipped_with_the_reason(tmp_path, monkeypatch):
    (tmp_path / 'paper.pdf').write_bytes(b'%PDF-1.4 not really')
    monkeypatch.setitem(sys.modules, 'pypdf', None)  # import fails
    monkeypatch.setattr(acad.shutil, 'which', lambda name: None)
    profile = acad.learn([tmp_path], tmp_path / 'out')
    assert profile['documents'] == 0 and 'pypdf' in profile['skipped'][0]


def test_reference_profile_holds_numbers_and_generic_phrases_only():
    ref = acad.reference_profile()
    assert sum(ref['corpus'].values()) >= 30 and ref['total_words'] > 100000
    assert all(s['doi'] and s['license'] for s in ref['sources'])
    phrases = [p for items in ref['phrasebank'].values() for p in items]
    assert phrases and all(len(p.split()) <= 4 for p in phrases)  # never sentences from the papers
    text = json.dumps(ref)
    assert len(text) < 40000
    assert 'Measured in high-impact journals' in acad.card('methods')
    assert ref['rates']['verbs_per_10k']['showed'] > ref['rates']['verbs_per_10k']['demonstrated']


def test_new_prose_rules_and_their_severity():
    def codes(text, section='discussion'):
        return {(sev, code) for sev, code, _, _ in acad.prose_issues(text, section)}
    assert ('high', 'ING_TAIL') in codes('Pain fell, highlighting the importance of early care.')
    assert codes('Pain fell, highlighting a gap.') == {('medium', 'ING_TAIL')}
    assert ('medium', 'INFLATION') in codes('This trial marks a pivotal milestone in spine care.')
    assert ('medium', 'VAGUE_ATTRIBUTION') in codes('Many believe that fusion is overused.')
    assert codes('Many believe that fusion is overused [EVID:a_2020].') == set()
    assert ('medium', 'SYNONYM_CYCLING') in codes('Surgeons utilize drains, leverage navigation and employ robots.')
    flat = ' '.join(['The cohort included older adults with stenosis today.'] * 5)
    assert ('medium', 'FLAT_RHYTHM') in codes(flat)
    assert ('medium', 'FLAT_RHYTHM') not in codes(flat, 'methods')


def test_preserve_flags_changed_numbers_citations_and_references(tmp_path):
    before = 'Pain fell (OR 2.3; 95% CI 1.2-4.5; *p* = 0.02) [EVID:a_2020] (Table 2).'
    assert acad.preserve_problems(before, 'Pain decreased (OR 2.3; 95% CI 1.2-4.5; *p* = 0.02) [EVID:a_2020] (Table 2).') == []
    problems = acad.preserve_problems(before, 'Pain fell (OR 2.3; 95% CI 1.2-4.6) [EVID:b_2021] (Table 3).')
    assert {"dropped '4.5'", "added '4.6'", "dropped '[EVID:a_2020]'", "dropped 'Table2'"} <= set(problems)
    (tmp_path / 'a.md').write_text(before, encoding='utf-8')
    (tmp_path / 'b.md').write_text(before.replace('2.3', '2.4'), encoding='utf-8')
    assert acad.main(['preserve', str(tmp_path / 'a.md'), str(tmp_path / 'b.md')]) == 1


def test_chat_toggles_the_mode_and_paper_folders_get_a_per_turn_reminder(tmp_path):
    intent = load_hook('style_intent')
    assert "now 'off'" in intent.evaluate({'prompt': '학술 모드 꺼줘'}) and acad.mode() == 'off'
    assert "now 'strict'" in intent.evaluate({'prompt': 'academic mode strict please'}) and acad.mode() == 'strict'
    assert "now 'academic'" in intent.evaluate({'prompt': '학술 모드 다시 켜줘'}) and acad.mode() == 'academic'
    (tmp_path / 'paper' / 'drafts').mkdir(parents=True)
    assert intent.evaluate({'prompt': '커밋해줘', 'cwd': str(tmp_path / 'paper')}) == ''  # a bare drafts/ is not a paper
    (tmp_path / 'paper' / 'project.json').write_text('{"artifacts": []}', encoding='utf-8')
    assert '[academic writing mode: academic]' in intent.evaluate({'prompt': '커밋해줘', 'cwd': str(tmp_path / 'paper')})
    assert intent.evaluate({'prompt': '커밋해줘', 'cwd': str(tmp_path)}) == ''
    acad.set_mode('off')
    assert intent.evaluate({'prompt': '커밋해줘', 'cwd': str(tmp_path / 'paper')}) == ''


def test_subagents_get_the_core_card_unless_off(tmp_path, monkeypatch):
    script = ROOT / 'scripts' / 'hooks' / 'subagent_style.py'
    out = subprocess.run([sys.executable, str(script)], input=json.dumps({'cwd': str(ROOT)}), capture_output=True,
                         text=True, encoding='utf-8').stdout  # the template checkout counts as a paper folder
    assert subprocess.run([sys.executable, str(script)], input=json.dumps({'cwd': str(tmp_path)}),
                          capture_output=True, text=True).stdout == ''
    payload = json.loads(out)['hookSpecificOutput']
    assert payload['hookEventName'] == 'SubagentStart' and 'ACADEMIC WRITING MODE' in payload['additionalContext']
    monkeypatch.setenv('MANUWRIGHT_WRITING_MODE', 'off')
    assert subprocess.run([sys.executable, str(script)], input='{}', capture_output=True, text=True).stdout == ''
