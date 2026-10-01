"""Synthetic whole-project tests for shared profiles, provenance and packaging."""
import json
from pathlib import Path
from unittest.mock import patch
import pytest
from harness.project import load_project, snapshot, verify, SEMANTIC_CHECKS
from harness.results import digest, validate_bindings
from harness.build import build
from harness.__main__ import main
from tests.test_enforce_gates import DRAFT_CONTENT, ANALYSIS_CONTENT


def put(path,value):
    path.parent.mkdir(parents=True,exist_ok=True)
    path.write_text(json.dumps(value,indent=2) if isinstance(value,dict) else value,encoding='utf-8')
    return path


@pytest.fixture(autouse=True)
def isolated_home(tmp_path_factory,monkeypatch):
    # The saved DOCX default lives in the user's config; never read the real one.
    home=tmp_path_factory.mktemp('mwhome');monkeypatch.setenv('MANUWRIGHT_HOME',str(home))
    return home


@pytest.fixture
def project(tmp_path):
    put(tmp_path/'drafts/03_introduction.md','# Introduction\n\nAn important clinical question remains unresolved.\n')
    put(tmp_path/'drafts/05_results.md','# Results\n\nMean age was 54 years.\n')
    put(tmp_path/'knowledge/evidence.md','# Evidence\n')
    put(tmp_path/'results/table.csv','age\n54\n')
    draft=put(tmp_path/'drafts/draft_plan.md',DRAFT_CONTENT+'\n- [x] 사용자 승인 완료\n')
    analysis=put(tmp_path/'data/analysis_plan.md',ANALYSIS_CONTENT+'\n- [x] 사용자 승인 완료\n')
    for plan in (draft,analysis):
        put(plan.with_suffix('.approval.json'),{'status':'approved','approved_by':'Synthetic reviewer',
            'decision_reference':'Synthetic test decision','sha256':digest(plan)})
    art=tmp_path/'drafts/05_results.md';csv=tmp_path/'results/table.csv'
    context=dict(outcome='age',timepoint='baseline',population='all',comparison='none',statistic='mean',unit='years')
    put(tmp_path/'review/bindings.json',{
        'results':[dict(result_id='p1.age',paper_id='p1',raw='54',source={'file':'results/table.csv',
            'row':2,'column':'age','sha256':digest(csv)},**context)],
        'bindings':[{'artifact':'drafts/05_results.md','artifact_sha256':digest(art),'token_index':0,
                     'result_id':'p1.age','context':context}]})
    put(tmp_path/'review/ai.json',{'used':False,'reviewed_by':'Synthetic reviewer'})
    put(tmp_path/'review/checklist.json',{'guideline':'Synthetic checklist','version':'1',
        'source_url':'https://example.org/checklist','reviewed_by':'Synthetic reviewer',
        'items':[{'id':'1','status':'PASS','location':'Results'}]})
    config={'schema_version':1,'paper_id':'p1','study_type':'original_research',
        'artifacts':['drafts/03_introduction.md','drafts/05_results.md'],'tables':[],
        'evidence':'knowledge/evidence.md','draft_plan':'drafts/draft_plan.md','analysis_plan':'data/analysis_plan.md',
        'results':'results','numeric_artifacts':['drafts/05_results.md'],'result_bindings':'review/bindings.json',
        'semantic_review':'review/semantic.json','human_signoff':'review/human.json',
        'ai_usage':'review/ai.json','checklist':'review/checklist.json'}
    return put(tmp_path/'project.json',config)


def sign(project):
    path,c=load_project(project);deps=snapshot(path,c)
    put(path.parent/'review/semantic.json',{'status':'PASS','reviewer':'Independent synthetic reviewer',
        'method':'independent','checks':dict.fromkeys(SEMANTIC_CHECKS,'PASS'),'findings':[], 'dependencies':deps})
    put(path.parent/'review/human.json',{'status':'approved','reviewer':'Synthetic human',
        'decision_reference':'Test only', 'dependencies':deps})


def test_draft_passes_submission_blocks_missing_reviews(project):
    assert verify(project)['status']=='PASS'
    assert verify(project,'submission')['status']=='BLOCKED'
    with pytest.raises(ValueError,match='submission verification'):
        build(project)
    assert not (project.parent/'output').exists()


def test_build_requires_fresh_reviews_and_preserves_sources(project):
    sign(project)
    report=verify(project,'submission')
    assert report['status']=='PASS',report
    before=digest(project.parent/'drafts/05_results.md')
    first=build(project);second=build(project)
    assert first!=second
    assert list(first.glob('manuscript_*.docx'))
    assert digest(project.parent/'drafts/05_results.md')==before
    from docx import Document
    doc=Document(next(first.glob('manuscript_*.docx')))
    assert 'Mean age was 54 years.' in '\n'.join(p.text for p in doc.paragraphs)
    assert all(not p.style.name.startswith('Heading') for p in doc.paragraphs)
    put(project.parent/'knowledge/evidence.md','# Evidence\nChanged\n')
    assert verify(project,'submission')['status']=='FAIL'


@pytest.mark.parametrize('key,value',[('outcome','fusion'),('timepoint','month12'),('unit','percent'),('comparison','control_minus_treatment')])
def test_context_swap_is_rejected(project,key,value):
    path=project.parent/'review/bindings.json';data=json.loads(path.read_text(encoding='utf-8'))
    data['bindings'][0]['context'][key]=value;put(path,data)
    assert verify(project)['status']=='BLOCKED'


def test_wrong_paper_rejected(project):
    path=project.parent/'review/bindings.json';data=json.loads(path.read_text(encoding='utf-8'))
    data['results'][0]['paper_id']='p2';put(path,data)
    assert verify(project)['status']=='BLOCKED'


def test_csv_change_invalidates_binding(project):
    put(project.parent/'results/table.csv','age\n99\n')
    assert verify(project)['status']!='PASS'


def test_analysis_plan_change_invalidates_approval(project):
    plan=project.parent/'data/analysis_plan.md';put(plan,plan.read_text(encoding='utf-8')+'\nChanged endpoint.\n')
    assert verify(project)['status']=='FAIL'


def test_receipt_does_not_claim_self_review_independence(project):
    sign(project);path=project.parent/'review/semantic.json';data=json.loads(path.read_text(encoding='utf-8'))
    data['method']='self';put(path,data)
    assert verify(project,'submission')['status']=='FAIL'


def test_dependency_added_after_review_is_stale(project):
    sign(project);put(project.parent/'results/new.csv','n\n80\n')
    assert verify(project,'submission')['status']=='FAIL'


def test_path_escape_rejected(project):
    data=json.loads(project.read_text(encoding='utf-8'));data['artifacts']=['../outside.md'];put(project,data)
    with pytest.raises(ValueError,match='escapes'):
        load_project(project)


def test_no_numerical_exemption_for_original_research(project):
    data=json.loads(project.read_text(encoding='utf-8'));data['numeric_artifacts']=[];data['numbers_not_applicable']='No numbers'
    put(project,data)
    report=verify(project)
    assert report['status']=='FAIL'
    assert any(c['check']=='numeric_scope' and c['status']=='FAIL' for c in report['checks'])


def test_packet_and_status_cli_share_snapshot(project,capsys):
    with patch('sys.argv',['harness','packet','--project',str(project)]):
        assert main()==0
    result=json.loads(capsys.readouterr().out)
    packet=json.loads(Path(result['packet']).read_text(encoding='utf-8'))
    assert packet['paper_id']=='p1'
    assert packet['files']
    with patch('sys.argv',['harness','verify','--project',str(project)]):
        assert main()==0
    capsys.readouterr()
    with patch('sys.argv',['harness','status','--project',str(project)]):
        assert main()==0
    assert json.loads(capsys.readouterr().out)['last_run']['status']=='PASS'


def test_changed_during_build_does_not_publish(project):
    sign(project)
    from harness import build as module
    original=module.append_markdown
    def mutate(doc,text,*style):
        original(doc,text,*style)
        put(project.parent/'knowledge/evidence.md','# Changed during build')
    with patch.object(module,'append_markdown',side_effect=mutate):
        with pytest.raises(ValueError,match='changed during build'):
            build(project)
    assert not list((project.parent/'output').glob('*'))

def test_build_stamps_revision_suffix(project):
    root=project.parent
    put(root/'drafts/revision/REV1/03_introduction_REV1.md',
        '# Introduction\n\nAn important clinical question remains unresolved. Additional context was added.\n')
    put(root/'review/comments_REV1.md','Reviewer #1:\n\nComment 1) Add context.\n')
    put(root/'drafts/revision/REV1/response_letter_REV1.md',
        '# Responses\n\nReviewer #1:\n\nComment 1) Add context.\n\n'
        '[CHANGE]\ncomment_id: R1-C1\nclaim: Added context\nsection: 03_introduction\n'
        'expected_terms: additional context\n[/CHANGE]\n\n'
        'Response: We thank the reviewer. Additional context was added.\n')
    data=json.loads(project.read_text(encoding='utf-8'))
    data['artifacts']=['drafts/revision/REV1/03_introduction_REV1.md','drafts/05_results.md']
    data['response']='drafts/revision/REV1/response_letter_REV1.md';data['comments']='review/comments_REV1.md'
    put(project,data);sign(project)
    assert verify(project,'revision')['status']=='PASS'
    out=build(project)
    names={item.name for item in out.iterdir()}
    assert any(name.startswith('manuscript_REV1_') and name.endswith('.docx') for name in names),names
    assert any(name.startswith('response_letter_REV1_') and name.endswith('.docx') for name in names),names
    assert not any(name.startswith('manuscript_2') for name in names),names


def test_revision_old_artifact_blocks_build(project):
    test_build_stamps_revision_suffix(project)
    data=json.loads(project.read_text(encoding='utf-8'))
    data['artifacts'][0]='drafts/03_introduction.md'
    put(project,data);sign(project)
    report=verify(project,'submission')
    assert any(c['check']=='revision_scope' and c['status']=='FAIL' for c in report['checks'])
    with pytest.raises(ValueError,match='revision_scope'):
        build(project)


def test_omitted_results_numbers_block_submission(project):
    data=json.loads(project.read_text(encoding='utf-8'))
    data['numeric_artifacts']=['drafts/03_introduction.md']
    put(project,data)
    put(project.parent/'review/bindings.json',{'results':[], 'bindings':[]})
    put(project.parent/'drafts/05_results.md','# Results\n\nMean age was 999 years.\n')
    sign(project)
    report=verify(project,'submission')
    assert any(c['check']=='numeric_scope' and c['status']=='FAIL' for c in report['checks'])
    with pytest.raises(ValueError,match='numeric_scope'):
        build(project)


@pytest.mark.parametrize('reason', ['', 'These are result values'])
def test_results_cannot_be_exempted(project, reason):
    data=json.loads(project.read_text(encoding='utf-8'))
    data['numeric_artifacts']=[]
    data['numeric_exemptions']={'drafts/05_results.md':reason}
    put(project,data)
    assert verify(project)['status']=='FAIL'


def test_nonresult_numeric_exemption_requires_reason(project):
    put(project.parent/'drafts/03_introduction.md','# Introduction\n\nPrior studies enrolled 80 participants.\n')
    data=json.loads(project.read_text(encoding='utf-8'))
    assert verify(project)['status']=='FAIL'
    data['numeric_exemptions']={'drafts/03_introduction.md':'Literature sample size; verify against cited evidence in semantic review.'}
    put(project,data)
    assert verify(project)['status']=='PASS'


def checklist_status(project, items):
    put(project.parent/'review/checklist.json',{'guideline':'Synthetic checklist','version':'1',
        'source_url':'https://example.org/checklist','reviewed_by':'Synthetic reviewer','items':items})
    sign(project)
    return next(c for c in verify(project,'submission')['checks'] if c['check']=='reporting_checklist')


@pytest.mark.parametrize('items', [[{'status':'PASS'}], [{'id':'1','status':'PASS'}],
    [{'id':'1','status':'PASS','location':'Results'},{'id':'1','status':'PASS','location':'Methods'}],
    [{'id':'2','status':'NOT_APPLICABLE'}], [{'id':'3','status':'TODO','location':'x'}], ['PASS']])
def test_incomplete_checklist_records_fail(project, items):
    assert checklist_status(project, items)['status']=='FAIL'


def test_complete_checklist_record_passes(project):
    assert checklist_status(project, [{'id':'1','status':'PASS','location':'Results'},
        {'id':'2','status':'NOT_APPLICABLE','reason':'No harms outcome'}])['status']=='PASS'


def test_used_ai_requires_tool_records(project):
    put(project.parent/'review/ai.json',{'used':True,'reviewed_by':'Synthetic reviewer',
        'disclosure':'Drafting assistance','tools':['some model']})
    sign(project)
    check=next(c for c in verify(project,'submission')['checks'] if c['check']=='ai_disclosure')
    assert check['status']=='FAIL' and 'tool and role' in check['detail']


def test_abstract_must_be_published_and_declared(project):
    put(project.parent/'drafts/02_abstract.md','# Abstract\n\nMean age was 54 years.\n')
    data=json.loads(project.read_text(encoding='utf-8'))
    data['abstract']='drafts/02_abstract.md'
    put(project,data)
    with pytest.raises(ValueError,match='published artifacts'):
        load_project(project)
    data['artifacts'].insert(0,'drafts/02_abstract.md'); data['numeric_artifacts'].append('drafts/02_abstract.md')
    del data['abstract']
    put(project,data)
    with pytest.raises(ValueError,match='declare it'):
        load_project(project)


def test_project_terminology_drives_lint_and_invalidates_review(project):
    put(project.parent/'Style/terminology.md','| Preferred Term | Forbidden Terms |\n|---|---|\n| issue | question |\n')
    data=json.loads(project.read_text(encoding='utf-8'))
    before=snapshot(*load_project(project))
    data['terminology']='Style/terminology.md'
    put(project,data)
    path,c=load_project(project)
    assert str(Path('Style/terminology.md')) in snapshot(path,c) and snapshot(path,c)!=before
    lint=next(x for x in verify(project)['checks'] if x['check']=='manuscript_lint')
    assert lint['status']=='FAIL' and '03_introduction.md' in lint['detail'], lint


def test_failed_citation_reports_location(project):
    put(project.parent/'drafts/03_introduction.md','# Introduction\n\nA claim [EVID:ghost_2020].\n')
    check=next(c for c in verify(project)['checks'] if c['check']=='citations')
    assert check['status']=='FAIL' and 'ghost_2020' in check['detail'] and 'line=3' in check['detail']


def test_packet_lists_omitted_sources(project,capsys):
    with patch('sys.argv',['harness','packet','--project',str(project)]):
        assert main()==0
    result=json.loads(capsys.readouterr().out)
    assert str(Path('results/table.csv')) in result['omitted_sources']
    assert str(Path('review/checklist.json')) not in result['omitted_sources']


def test_project_style_spec_is_checked(project):
    put(project.parent/'drafts/style_spec.md','| Section | Words | Mean sentence | Paragraphs |\n|---|---|---|---|\n| results | 400 | 20 | 4 |\n')
    data=json.loads(project.read_text(encoding='utf-8'))
    data['style_spec']='drafts/style_spec.md'
    put(project,data)
    check=next(c for c in verify(project)['checks'] if c['check']=='style_metrics')
    assert check['status']=='FAIL' and '05_results.md' in check['detail'], check


def test_draft_runs_before_submission_records_exist(project):
    (project.parent/'review/ai.json').unlink(); (project.parent/'review/checklist.json').unlink()
    assert verify(project)['status']=='PASS'
    with patch('sys.argv',['harness','packet','--project',str(project)]):
        assert main()==0  # packet must not require records that do not exist yet
    sign(project)
    report=verify(project,'submission')
    assert report['status']=='BLOCKED'
    assert {c['check'] for c in report['checks'] if c['status']=='BLOCKED'} >= {'ai_disclosure','reporting_checklist'}


def test_docx_style_user_default_then_project_override(project,isolated_home):
    from harness.build import docx_style
    from docx import Document
    (isolated_home/'config.json').write_text(json.dumps({'docx':{'font':'Arial','size':11,'line_numbers':'page'}}))
    config=json.loads(project.read_text());config['docx']={'size':12,'line_spacing':1.5,'page_numbers':'right'}
    project.write_text(json.dumps(config))
    style=docx_style(config)
    assert (style['font'],style['size'],style['line_spacing'],style['line_numbers'])==('Arial',12,1.5,'page')
    sign(project)
    out=build(project)
    doc=Document(next(out.glob('manuscript_*.docx')))
    normal=doc.styles['Normal']
    assert normal.font.name=='Arial' and normal.font.size.pt==12 and normal.paragraph_format.line_spacing==1.5
    assert doc.sections[0]._sectPr.xpath('./w:lnNumType')[0].get(
        '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}restart')=='newPage'
    assert doc.sections[0].footer.paragraphs[0].alignment==2
    assert json.loads((out/'build.json').read_text())['docx_style']['font']=='Arial'
    for bad in ({'colour':'red'},{'size':'big'},{'line_numbers':'sometimes'}):
        with pytest.raises(ValueError,match='docx'):
            docx_style({'docx':bad})


def test_build_uses_journal_preset(project):
    from docx import Document
    put(project.parent/'knowledge/evidence.md','# Evidence\n\n### [1] Zed 2020\n- **Evidence ID:** zed_2020\n'
        '- **Citation:** Zed A, Young B. Later trial. Spine. 2020;45(1):1-9.\n- **Source Status:** verified\n')
    put(project.parent/'drafts/03_introduction.md','# Introduction\n\nAn important clinical question remains unresolved [EVID:zed_2020].\n')
    config=json.loads(project.read_text());config['journal']='ama';put(project,config)
    sign(project)
    out=build(project)
    doc=Document(next(out.glob('manuscript_*.docx')))
    assert any(run.font.superscript and run.text=='1' for p in doc.paragraphs for run in p.runs)
    assert any('Zed A, Young B. Later trial. *Spine*' in p.text or 'Zed A, Young B. Later trial. Spine' in p.text for p in doc.paragraphs)
    assert json.loads((out/'build.json').read_text())['citation_style']=='ama preset'
    config['journal']='no-such-journal';put(project,config)
    with pytest.raises(ValueError,match='unknown journal preset'):
        load_project(project)


def test_chat_approval_ticks_box_and_writes_receipt(tmp_path):
    from harness.__main__ import approve_in_chat
    from harness.project import checker
    plan=put(tmp_path/'data/analysis_plan.md',ANALYSIS_CONTENT+'\n- [ ] 사용자 승인 완료\n')
    data=approve_in_chat(plan,'analysis','Dr. Author','승인')
    text=plan.read_text(encoding='utf-8')
    assert '- [x] 사용자 승인 완료 — Dr. Author' in text and '채팅 승인: "승인"' in text and '- [ ]' not in text
    assert checker('plan_validation').approval_problem(plan,required=True) is None
    assert data['decision_reference']=='chat approval: "승인"'
    put(plan,plan.read_text(encoding='utf-8')+'\nchanged after approval\n')
    assert 'stale' in checker('plan_validation').approval_problem(plan,required=True)
    with pytest.raises(ValueError,match='incomplete'):
        approve_in_chat(put(tmp_path/'p2.md','# Analysis Plan\n'),'analysis','Dr. Author','승인')
    with pytest.raises(ValueError,match='blank'):
        approve_in_chat(put(tmp_path/'p3.md',ANALYSIS_CONTENT),'analysis','Dr. Author',' ')
