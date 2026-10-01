"""Manifest-ordered DOCX and Markdown package, gated before publication."""
from __future__ import annotations
import json
import os
import re
import shutil
import tempfile
from datetime import datetime, timezone
from pathlib import Path
from uuid import uuid4
from .project import load_project, verify, snapshot, checker
from .results import inside, digest


# Manuscript DOCX style. Precedence: these defaults < the user's saved default
# (`manuwright config set docx.<key> <value>`, ~/.manuwright/config.json) < project.json "docx"
# (per paper, i.e. per target journal).
DOCX_DEFAULTS = {'font': 'Times New Roman', 'size': 10, 'heading_size': 12, 'subheading_size': 11,
                 'line_spacing': 2.0, 'margin_inches': 1.0,
                 'line_numbers': 'continuous', 'page_numbers': 'center'}
DOCX_CHOICES = {'line_numbers': ('continuous', 'page', 'off'), 'page_numbers': ('center', 'right', 'off')}


def user_docx():
    try:
        home = Path(os.environ.get('MANUWRIGHT_HOME') or Path.home() / '.manuwright')
        return json.loads((home / 'config.json').read_text(encoding='utf-8')).get('docx', {})
    except (OSError, ValueError, RuntimeError):  # no saved config, or no home directory
        return {}


def check_docx_values(source, values, allow_reference=False):
    if not isinstance(values, dict):
        raise ValueError(f'{source} docx must be an object')
    for key, value in values.items():
        if key == 'reference' and allow_reference:
            if not (isinstance(value, str) and value.lower().endswith('.docx')):
                raise ValueError(f'{source} docx.reference must be a .docx path inside the paper folder')
            continue
        if key not in DOCX_DEFAULTS:
            raise ValueError(f'{source} docx: unknown key {key!r} (known: {", ".join(DOCX_DEFAULTS)}, reference)')
        if key in DOCX_CHOICES and value not in DOCX_CHOICES[key]:
            raise ValueError(f'{source} docx.{key} must be one of {", ".join(DOCX_CHOICES[key])}')
        if key == 'font' and not (isinstance(value, str) and value.strip()):
            raise ValueError(f'{source} docx.font must be a font name')
        if key not in DOCX_CHOICES and key != 'font' and not (
                isinstance(value, (int, float)) and not isinstance(value, bool) and 0 < value <= 72):
            raise ValueError(f'{source} docx.{key} must be a positive number')


def docx_style(config, root=None):
    """Resolved Word style. With "reference" (a .docx designed in Word) the template's own styles
    are the base and only keys written in project.json are applied on top; otherwise defaults <
    user default (~/.manuwright/config.json) < project.json."""
    project = config.get('docx', {})
    check_docx_values('project.json', project, allow_reference=True)
    if isinstance(project, dict) and project.get('reference'):
        if root is None:
            raise ValueError('docx.reference needs the paper folder')
        path = inside(root, project['reference'])  # refuses paths that leave the paper folder
        if not path.is_file():
            raise ValueError(f'docx.reference not found: {project["reference"]}')
        try:
            from docx import Document
            Document(str(path))
        except ImportError:
            pass
        except Exception as exc:  # renamed .dotx, corrupt file: fail before verification, not after
            raise ValueError(f'docx.reference is not a Word document python-docx can open: {exc}') from exc
        return {**{k: v for k, v in project.items() if k != 'reference'}, 'reference': str(path),
                'reference_path': project['reference']}
    style = dict(DOCX_DEFAULTS)
    check_docx_values('user config', user_docx())
    style.update(user_docx())
    style.update(project)
    return style


def document(style=DOCX_DEFAULTS, numbered=True):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    if style.get('reference'):
        doc = Document(style['reference'])  # the author's Word template: keep its styles, drop its text
        body = doc.element.body
        for child in list(body):
            if not child.tag.endswith('}sectPr'):
                body.remove(child)
    else:
        doc = Document()
    normal = doc.styles['Normal']
    if 'font' in style:
        normal.font.name = style['font']
    if 'size' in style:
        normal.font.size = Pt(style['size'])
    if not style.get('reference'):
        normal.font.color.rgb = RGBColor(0, 0, 0)
    if 'line_spacing' in style:
        normal.paragraph_format.line_spacing = style['line_spacing']
    def has_page_field(footer):  # anywhere in the footer, including Word's content controls (w:sdt)
        root = footer._element
        return any('PAGE' in (node.get(qn('w:instr')) or '') for node in root.iter(qn('w:fldSimple'))) or \
            any('PAGE' in (node.text or '') for node in root.iter(qn('w:instrText')))

    for section in doc.sections:
        if 'margin_inches' in style:
            section.top_margin = section.bottom_margin = Inches(style['margin_inches'])
            section.left_margin = section.right_margin = Inches(style['margin_inches'])
        sect = section._sectPr
        if numbered and 'line_numbers' in style:
            for old in sect.findall(qn('w:lnNumType')):  # a template may already number lines
                sect.remove(old)
            if style['line_numbers'] != 'off':
                lines = OxmlElement('w:lnNumType')
                lines.set(qn('w:countBy'), '1')
                lines.set(qn('w:restart'), 'newPage' if style['line_numbers'] == 'page' else 'continuous')
                # schema order: lnNumType comes before pgNumType, cols, ... docGrid
                after = [sect.find(qn('w:' + tag)) for tag in ('pgNumType', 'cols', 'formProt', 'vAlign', 'noEndnote',
                                                              'titlePg', 'textDirection', 'bidi', 'rtlGutter', 'docGrid')]
                after = [node for node in after if node is not None]
                after[0].addprevious(lines) if after else sect.append(lines)
        if numbered and style.get('page_numbers', 'off') != 'off':
            paragraph = section.footer.paragraphs[0]
            if not has_page_field(section.footer):  # keep a template's own page number
                paragraph.alignment = 2 if style['page_numbers'] == 'right' else 1
                field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE')
                paragraph._p.append(field)
    return doc


def inline(paragraph, text):
    # Intentionally bounded Markdown subset; do not silently turn complex
    # math/images/code into inaccurate Word output.
    for part in re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*|\^[^^\s]+\^)', text):
        if not part:
            continue
        if part.startswith('^') and part.endswith('^') and len(part) > 2:  # journal superscript citation
            paragraph.add_run(part[1:-1]).font.superscript = True
        elif part.startswith('**') and part.endswith('**'):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith('*') and part.endswith('*'):
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def append_markdown(doc, text, style=DOCX_DEFAULTS):
    from docx.shared import Pt
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    text = re.sub(r'<!--.*?-->', '', text, flags=re.S)
    if '```' in text or '$$' in text or '![' in text:
        raise ValueError('DOCX compiler supports prose/headings/pipe tables; move code, math or embedded images to separately reviewed supplements')
    lines = text.splitlines(); i = 0
    while i < len(lines):
        line = lines[i].strip(); i += 1
        if not line or line == '---':
            continue
        if line.startswith('|'):
            rows = [line]
            while i < len(lines) and lines[i].strip().startswith('|'):
                rows.append(lines[i].strip()); i += 1
            cells = [[cell.strip() for cell in row.strip('|').split('|')] for row in rows
                     if not re.fullmatch(r'[|\s:\-]+', row)]
            if not cells or any(len(row) != len(cells[0]) for row in cells):
                raise ValueError('malformed Markdown table')
            table = doc.add_table(rows=len(cells), cols=len(cells[0]))
            borders = OxmlElement('w:tblBorders')
            for side in ('top','bottom','left','right','insideH','insideV'):
                element = OxmlElement('w:'+side); element.set(qn('w:val'), 'nil'); borders.append(element)
            table._tbl.tblPr.append(borders)
            for row_index, row in enumerate(cells):
                for column, value in enumerate(row):
                    cell = table.cell(row_index,column); inline(cell.paragraphs[0],value)
                    if row_index == 0:
                        for run in cell.paragraphs[0].runs: run.bold=True
                    edges = OxmlElement('w:tcBorders')
                    for side in (('top','bottom') if row_index == 0 else ('bottom',) if row_index == len(cells)-1 else ()):
                        edge=OxmlElement('w:'+side);edge.set(qn('w:val'),'single');edge.set(qn('w:sz'),'6');edges.append(edge)
                    cell._tc.get_or_add_tcPr().append(edges)
            continue
        match = re.match(r'^(#{1,6})\s+(.+)',line)
        if match:
            level=min(len(match.group(1)),2)
            if style.get('reference') and f'Heading {level}' in [st.name for st in doc.styles]:
                doc.add_paragraph(match.group(2),style=f'Heading {level}')  # the template's own heading style
                continue
            run=doc.add_paragraph().add_run(match.group(2));run.bold=True
            run.italic=level>1
            size=style.get('heading_size' if level==1 else 'subheading_size')
            if size: run.font.size=Pt(size)
        else:
            prose=[line]
            while i<len(lines) and lines[i].strip() and not lines[i].lstrip().startswith(('#','|','- ')):
                prose.append(lines[i].strip());i+=1
            paragraph=doc.add_paragraph();inline(paragraph,' '.join(prose))


def build(path):
    path, config = load_project(path);root=path.parent
    style=docx_style(config,root)  # fail on a bad style before the (slower) verification
    report=verify(path,'submission')
    if report['status']!='PASS':
        raise ValueError('submission verification '+report['status']+': '+', '.join(x['check'] for x in report['checks'] if x['status'] in {'FAIL','BLOCKED'}))
    # Also freeze review receipts, which are not part of their own dependency map.
    receipts = {key:digest(inside(root,config[key])) for key in ('semantic_review','human_signoff')}
    formatter=checker('format_references')
    paths=[inside(root,item) for item in config['artifacts']+config.get('tables',[])+config.get('supplements',[])]
    journal=config.get('journal')  # a journal_styles preset: reference format + in-text markers
    references=formatter.build(paths,evidence_path=inside(root,config['evidence']),style='numbered',journal=journal)
    if references.unknown or references.missing_citation or (journal and (references.incomplete_authors or references.unparsed)):
        raise ValueError('incomplete bibliography'+(': '+', '.join(references.incomplete_authors+references.unparsed) if journal else ''))
    marker=references.in_text if journal else None
    date=datetime.now().strftime('%y%m%d')
    # Rule 5: revision packages carry _REVn before the date (manuscript_REV1_YYMMDD.docx).
    stamp='_'+date
    if config.get('response'):
        folder=inside(root,config['response']).parent
        revision=re.fullmatch(r'REV(\d+)',folder.name)
        if revision and folder.parent.name=='revision':
            stamp=f'_REV{revision.group(1)}_{date}'
    destination=inside(root, config.get('output','output'))
    destination.mkdir(parents=True,exist_ok=True)
    final=destination/(datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%SZ')+'_'+uuid4().hex[:8])
    with tempfile.TemporaryDirectory(prefix='.building-',dir=destination) as tmp:
        stage=Path(tmp);manuscript=document(style);merged=[];section_count=0
        for value in config['artifacts']:
            source=inside(root,value)
            text=formatter.convert_text(source.read_text(encoding='utf-8'),references.labels,marker)
            if source.name.startswith('08_references'):
                continue  # bibliography is generated below from actual first appearances
            if source.name.startswith('01_title'):
                title=document(style,False);append_markdown(title,text,style);title.save(stage/f'title_page{stamp}.docx');continue
            if section_count: manuscript.add_page_break()
            append_markdown(manuscript,text,style);merged.append(text);section_count+=1
        if references.references:
            bibliography='# References\n\n'+'\n\n'.join(references.references)
            manuscript.add_page_break();append_markdown(manuscript,bibliography,style);merged.append(bibliography)
        if config.get('ai_usage'):
            ai=json.loads(inside(root,config['ai_usage']).read_text(encoding='utf-8'))
            if ai.get('used'):
                disclosure='# AI assistance disclosure\n\n'+ai['disclosure']
                manuscript.add_page_break();append_markdown(manuscript,disclosure,style);merged.append(disclosure)
        manuscript.save(stage/f'manuscript{stamp}.docx')
        (stage/f'manuscript{stamp}.md').write_text('\n\n'.join(merged),encoding='utf-8')
        for index,value in enumerate(config.get('tables',[]),1):
            text=formatter.convert_text(inside(root,value).read_text(encoding='utf-8'),references.labels,marker)
            table=document(style,False);append_markdown(table,text,style);table.save(stage/f'table_{checker("check_crossrefs").TABLE_FILE_RE.search(Path(value).stem).group(1)}{stamp}.docx')
        for value in config.get('supplements',[]):  # separate files, never numbered as main tables
            text=formatter.convert_text(inside(root,value).read_text(encoding='utf-8'),references.labels,marker)
            supplement=document(style,False);append_markdown(supplement,text,style)
            supplement.save(stage/f'supplementary_{re.sub(r"^supp(?:lementary)?[_-]?","",Path(value).stem,flags=re.I)}{stamp}.docx')
        for index,value in enumerate(config.get('figures',[]),1):
            source=inside(root,value);shutil.copyfile(source,stage/f'figure_{checker("check_crossrefs").FIGURE_FILE_RE.search(source.stem).group(1)}{source.suffix}')
        if config.get('response'):
            response=checker('compile_response_docx')
            source=inside(root,config['response'])
            response_refs=formatter.build([source],evidence_path=inside(root,config['evidence']),style='numbered')
            if response_refs.unknown or response_refs.missing_citation:
                raise ValueError('incomplete response bibliography')
            response_text=formatter.convert_text(source.read_text(encoding='utf-8'),response_refs.labels)
            if response_refs.references:
                response_text+='\n\nReferences\n\n'+'\n\n'.join(response_refs.references)
            temporary_response=stage/'response_source.md';temporary_response.write_text(response_text,encoding='utf-8')
            response.compile_docx(temporary_response,stage/f'response_letter{stamp}.docx',keep_change_markers=False)
            temporary_response.unlink()
        if snapshot(path,config)!=report['dependencies'] or any(digest(inside(root,config[key]))!=value for key,value in receipts.items()):
            raise ValueError('inputs or approvals changed during build')
        (stage/'verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
        outputs={item.name:digest(item) for item in sorted(stage.iterdir()) if item.is_file()}
        (stage/'build.json').write_text(json.dumps({'schema_version':1,'paper_id':config['paper_id'],
            'outputs':outputs,'visual_qa':'required before submission','citation_style':(journal+' preset' if journal else 'numbered; source Citation strings preserved'),'docx_style':{k:v for k,v in style.items() if k!='reference'}},indent=2),encoding='utf-8')
        stage.rename(final)
    return final
